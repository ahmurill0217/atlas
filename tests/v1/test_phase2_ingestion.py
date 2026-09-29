"""Phase 2: PDF / DOCX / generic JSON adapters and segmentation."""

import json
import uuid
from datetime import datetime

import docx
import pytest
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session

from atlas.compiler.compiler import GraphCompiler
from atlas.db.models import DocumentVersion, EdgeEvidence
from atlas.extraction.structured import StructuredExtractor
from atlas.graph.queries import GraphQueries
from atlas.ingestion.adapters import UnsupportedSource, normalize_file
from atlas.ingestion.segmentation import Block, assemble, headed_blocks, split_long
from atlas.ontology import load_ontology
from atlas.pipeline import KnowledgeIngestionPipeline, segment
from tests.v1.conftest import CORPUS, ROOT, candidate_set, edge, ent


@pytest.fixture(scope="session")
def ontology_v12():
    return load_ontology(str(ROOT / "ontology" / "v1_2"))


# --- segmentation -----------------------------------------------------------------

def test_split_long_keeps_sentences_whole():
    text = ("Northwind Inc. will deploy Atlas in the U.S. next month. " * 3 + "Short one. ") * 4
    parts = split_long(text, max_chars=200)
    assert all(len(p) <= 200 for p in parts) and len(parts) > 1
    assert all(p.endswith(".") for p in parts)                   # never cut mid-sentence
    assert " ".join(parts) == " ".join(text.split())
    assert split_long("x" * 50 + " " + "y" * 50, max_chars=60) == ["x" * 50, "y" * 50]   # no boundary: whitespace


def test_assemble_offsets_and_parts():
    raw, sections = assemble([Block("paragraph", "A. " * 400, {"page_number": 3}), Block("heading", "Title")],
                             max_chars=500)
    assert sections[0].metadata["parts"] == len([s for s in sections if s.kind == "paragraph"]) > 1
    assert all(s.metadata.get("page_number") == 3 for s in sections if s.kind == "paragraph")
    assert [s.ordinal for s in sections] == list(range(len(sections)))
    for s in sections:
        assert raw[s.start_char:s.end_char] == s.text


def test_heading_paths():
    blocks = headed_blocks([(0, "SOW"), (1, "Scope"), (None, "text a"), (2, "Timeline"), (None, "text b"),
                            (1, "Fees"), (None, "| table |", "table")])
    paths = {b.text: b.metadata["heading_path"] for b in blocks}
    assert paths["text b"] == ["SOW", "Scope", "Timeline"]
    assert paths["| table |"] == ["SOW", "Fees"] and blocks[-1].kind == "table"


# --- adapters -----------------------------------------------------------------------

def _pdf(path, pages, title="Quarterly Review", author="Microsoft Office User"):
    c = canvas.Canvas(str(path))
    c.setTitle(title)
    c.setAuthor(author)
    for lines in pages:
        y = 720
        for line in lines:
            c.drawString(72, y, line)
            y -= 16
        c.showPage()
    c.save()
    return path


def test_pdf_adapter_is_page_aware(tmp_path):
    path = _pdf(tmp_path / "r.pdf", [
        ["Overview", "Acme renewed the Atlas contract for another year of ser-", "vice and support.",
         "Tom Becker approved it."],
        ["Risks", "1. SSO rollout is late.", "2. Budget is not yet approved for the second phase of work."]])
    doc = segment(normalize_file(path))
    assert doc.source_type == "pdf" and doc.title == "Quarterly Review" and doc.metadata["pages"] == 2
    assert [(s.kind, s.metadata["page_number"]) for s in doc.sections] == [
        ("heading", 1), ("paragraph", 1), ("heading", 2), ("paragraph", 2), ("paragraph", 2)]
    # hyphenation repaired; a full-width line continues its paragraph
    assert doc.sections[1].text == "Acme renewed the Atlas contract for another year of service and support. Tom Becker approved it."
    assert doc.sections[3].text == "1. SSO rollout is late."                   # list items stay separate
    assert doc.participants == [] and doc.metadata["file_metadata_author"] == "Microsoft Office User"


def test_scanned_and_corrupt_pdfs(tmp_path):
    blank = _pdf(tmp_path / "scan.pdf", [[]])
    doc = normalize_file(blank)
    assert doc.sections == [] and doc.metadata["needs_ocr"] and doc.metadata["pages_without_text"] == [1]
    (tmp_path / "broken.pdf").write_bytes(b"%PDF-1.4 not really")
    with pytest.raises(UnsupportedSource):
        normalize_file(tmp_path / "broken.pdf")


def test_docx_adapter_reads_headings_and_tables_in_order(tmp_path):
    d = docx.Document()
    d.core_properties.title, d.core_properties.author = "Master Services Agreement", "Legal"
    d.core_properties.created = datetime(2026, 1, 2, 3, 4, 5)
    d.add_heading("Master Services Agreement", level=0)
    d.add_heading("Parties", level=1)
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Customer", "Acme Corp"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Vendor", "Northwind Inc."
    d.add_heading("Term", level=1)
    d.add_paragraph("The agreement runs for 24 months.")
    d.save(tmp_path / "msa.docx")
    doc = segment(normalize_file(tmp_path / "msa.docx"))
    assert [s.kind for s in doc.sections] == ["heading", "heading", "table", "heading", "paragraph"]
    assert doc.sections[2].text == "Customer | Acme Corp\nVendor | Northwind Inc."
    assert doc.sections[4].metadata["heading_path"] == ["Master Services Agreement", "Term"]
    assert doc.created_at.year == 2026 and doc.metadata["file_metadata_author"] == "Legal"


def test_generic_json_adapter(tmp_path):
    doc = normalize_file(CORPUS / "drive" / "atlas_pilot_plan.json")
    assert (doc.source_system, doc.source_type, doc.source_external_id) == ("gdrive", "google_doc", "1Qx7AtlasPilotPlan")
    assert [(p.role, p.email) for p in doc.participants] == [
        ("author", "priya.shah@northwind.io"), ("owner", "mike.rodriguez@northwind.io"),
        ("editor", "sarah.chen@acme.com"), ("viewer", None)]
    assert doc.sections[-1].kind == "table" and doc.sections[-1].metadata["heading_path"][-1] == "Workstreams"
    for bad in ({"atlas_document": "9.9", "source_system": "x", "id": "1", "text": "t"},
                {"atlas_document": "1.0", "source_system": "x", "text": "t"},
                {"atlas_document": "1.0", "source_system": "x", "id": "1"}):
        (tmp_path / "bad.json").write_text(json.dumps(bad))
        with pytest.raises(UnsupportedSource):
            normalize_file(tmp_path / "bad.json")


def test_document_facts_from_structured_metadata(ontology_v12):
    doc = normalize_file(CORPUS / "drive" / "atlas_pilot_plan.json")
    cs = StructuredExtractor(ontology_v12, ["northwind.io"]).extract(doc, uuid.uuid4())
    edges = {(e.source_local_id, e.suggested_relation, e.target_local_id) for e in cs.edges}
    assert ("document", "AUTHORED_BY", "person:email:priya.shah@northwind.io") in edges
    assert {e.suggested_relation for e in cs.edges} == {"AUTHORED_BY", "WORKS_AT"}           # roles = metadata only
    assert "person:email:mike.rodriguez@northwind.io" in {e.local_id for e in cs.entities}   # owner still resolved
    assert any(p.role == "owner" for p in doc.participants)                                  # role kept in record
    pdf = normalize_file(CORPUS / "docs" / "Atlas_Security_Review.pdf")
    pdf_cs = StructuredExtractor(ontology_v12).extract(pdf, uuid.uuid4())
    assert [e.suggested_type for e in pdf_cs.entities] == ["Document"] and pdf_cs.edges == []


# --- pipeline ------------------------------------------------------------------------

def test_full_corpus_under_v12(kg, atlas_settings, ontology_v12):
    report = KnowledgeIngestionPipeline(kg, atlas_settings, ontology_v12).ingest(CORPUS)
    assert report.stats["documents_processed"] == 9 and "documents_failed" not in report.stats
    with Session(kg) as s:
        q = GraphQueries(s)
        rels = {(g["source_name"], g["relation_type"], g["target_name"]) for g in q.get_edges()}
        assert ("Atlas pilot plan", "AUTHORED_BY", "Priya Shah") in rels
        assert not q.find_entity(name="Microsoft Office User")
        titles = {e["canonical_name"] for e in q.list_entities("Document")}
        assert {"Statement of Work: Atlas Rollout", "Atlas Security Review", "Atlas pilot plan"} <= titles
        assert q.list_reviews(review_type="NEW_ONTOLOGY_CANDIDATE") == []
        assert q.stats()["unsupported_edges"] == 0
    again = KnowledgeIngestionPipeline(kg, atlas_settings, ontology_v12).ingest(CORPUS)
    assert again.stats["documents_unchanged"] == 9


def test_text_evidence_records_pdf_page(kg, atlas_settings, ontology_v12):
    """Evidence that points at a section inherits the section's page number."""
    pipeline = KnowledgeIngestionPipeline(kg, atlas_settings, ontology_v12)
    pipeline.ingest(CORPUS / "docs" / "Atlas_Security_Review.pdf")
    doc = normalize_file(CORPUS / "docs" / "Atlas_Security_Review.pdf")
    with Session(kg) as s:
        vid = s.execute(select(DocumentVersion.id).where(DocumentVersion.document_id == doc.document_id)).scalar_one()
        finding = next(sec for sec in doc.sections if sec.text.startswith("1. SSO"))
        cand = edge(doc, vid, "tom", "acme", "WORKS_AT", "Tom Becker leads the review on the Acme side.")
        cand.evidence.section_ordinal = finding.ordinal
        report = GraphCompiler(s, ontology_v12).compile(doc, candidate_set(doc, vid, [
            ent("tom", "Person", "Tom Becker", email="tom.becker@acme.com"),
            ent("acme", "Organization", "Acme Corp", domain="acme.com")], [cand]))
        assert report.edges == {"tom|WORKS_AT|acme": "ACCEPTED"}
        ev = s.execute(select(EdgeEvidence).where(EdgeEvidence.section_id.is_not(None))).scalar_one()
        assert ev.page_number == 2

"""Regenerate the binary sample documents in examples/business/docs (DOCX, PDF).

    uv run python scripts/make_business_samples.py

Outputs are committed; the script exists so they are reproducible and reviewable.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import docx
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parents[1] / "examples" / "business" / "docs"


def make_sow(path: Path) -> None:
    d = docx.Document()
    d.core_properties.title = "Statement of Work: Atlas Rollout"
    d.core_properties.author = "Mike Rodriguez"
    d.core_properties.created = datetime(2026, 9, 10, 14, 0, 0)
    d.core_properties.modified = datetime(2026, 9, 12, 9, 30, 0)
    d.add_heading("Statement of Work: Atlas Rollout", level=0)
    d.add_heading("Parties", level=1)
    table = d.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, [("Party", "Role", "Contact"),
                                        ("Acme Corp", "Customer", "Sarah Chen"),
                                        ("Northwind Inc.", "Vendor", "Mike Rodriguez")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    d.add_heading("Scope", level=1)
    d.add_paragraph("Northwind Inc. will deploy the Atlas AI assistant for the Acme Corp Boston team. "
                    "Mike Rodriguez owns the implementation plan. The pilot covers approx. 40 users in the U.S. "
                    "and excludes the EU entities, e.g. Acme GmbH.")
    d.add_heading("Timeline", level=2)
    d.add_paragraph("The pilot starts on October 5, 2026 and must be live by November 30, 2026. "
                    "The security review by Acme's IT team is a prerequisite for go-live.")
    d.add_heading("Commercials", level=1)
    d.add_paragraph("The contract value is USD 120,000 for the pilot phase.")
    d.save(path)


def make_security_review(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=LETTER)
    c.setTitle("Atlas Security Review")
    c.setAuthor("Microsoft Office User")         # typical junk file-metadata author
    c.setCreator("Acme IT")

    def page(lines: list[str]) -> None:
        y = 720
        for line in lines:
            if line:
                c.drawString(72, y, line)
            y -= 16 if line else 28
        c.showPage()

    page(["Atlas Security Review", "",
          "Prepared by Acme Corp IT Security for the Atlas rollout with Northwind.",
          "Tom Becker leads the review on the Acme side.", "",
          "Summary: the Atlas assistant processes internal documents and meeting tran-",
          "scripts. Data residency and SSO integration are the main areas of concern."])
    page(["Findings", "",
          "1. SSO via Okta is required before the pilot can go live.",
          "2. Meeting transcripts must be retained for no more than 90 days.", "",
          "Status: conditionally approved, pending the SSO integration."])
    c.save()


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_sow(OUT / "Atlas_SOW.docx")
    make_security_review(OUT / "Atlas_Security_Review.pdf")
    print("wrote", sorted(p.name for p in OUT.iterdir()))
